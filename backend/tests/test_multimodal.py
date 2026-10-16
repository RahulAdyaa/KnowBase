"""Tests for multimodal extraction (PDF/PPTX/DOCX/image) and ingestion."""

import io
import os

os.environ["VECTOR_BACKEND"] = "memory"
os.environ["EMBEDDING_PROVIDER"] = "fake"
os.environ["DATABASE_URL"] = "sqlite://"
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-key-min-32-chars!!")

import pytest  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from app.core import security  # noqa: E402
from app.db.base import Base  # noqa: E402
from app.models.credit import CreditAccount, CreditLedger  # noqa: E402
from app.models.document import Document, DocumentChunk, DocumentStatus  # noqa: E402
from app.models.user import Role, User  # noqa: E402
from app.services.embeddings import FakeEmbeddingProvider  # noqa: E402
from app.services.extractors import extract_segments, supported_types  # noqa: E402
from app.services.rag import ingest_document  # noqa: E402
from app.services.vectorstore import InMemoryVectorStore  # noqa: E402

security.settings.jwt_secret_key = "test-secret-key-min-32-chars!!"


def _make_pdf_bytes(text: str) -> bytes:
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
    ]
    stream = b"BT /F1 24 Tf 100 700 Td (" + text.encode() + b") Tj ET"
    objs.append(b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream")
    objs.append(b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode())
    for o in offsets:
        out.write(f"{o:010d} 00000 n \n".encode())
    out.write(f"trailer<</Size {len(objs)+1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def _make_pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Quarterly revenue grew 12 percent"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Onboarding guide for new engineers")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_png_bytes() -> bytes:
    from PIL import Image

    img = Image.new("RGB", (64, 64), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_supported_types_cover_multimodal():
    types = supported_types()
    for ext in (".pdf", ".pptx", ".docx", ".txt", ".png", ".mp3", ".mp4"):
        assert ext in types


def test_extract_pdf():
    segs = extract_segments("memo.pdf", "application/pdf", _make_pdf_bytes("budget approved"))
    assert len(segs) == 1 and segs[0].page_number == 1
    assert "budget approved" in segs[0].text


def test_extract_pptx():
    segs = extract_segments("deck.pptx", "", _make_pptx_bytes())
    assert len(segs) == 1 and segs[0].page_number == 1
    assert "revenue" in segs[0].text


def test_extract_docx():
    segs = extract_segments("notes.docx", "", _make_docx_bytes())
    assert "Onboarding guide" in segs[0].text


def test_extract_image_modality():
    segs = extract_segments("chart.png", "image/png", _make_png_bytes())
    assert segs[0].modality == "image"


def _db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    tables = [User.__table__, CreditAccount.__table__, CreditLedger.__table__,
              Document.__table__, DocumentChunk.__table__]
    Base.metadata.create_all(engine, tables=tables)
    return sessionmaker(bind=engine, expire_on_commit=False)()


def test_ingest_pdf_end_to_end():
    session = _db()
    user = User(email="mm@example.com", hashed_password=security.hash_password("pw12345678"), role=Role.USER)
    session.add(user)
    session.commit()
    store, embedder = InMemoryVectorStore(), FakeEmbeddingProvider()
    raw = _make_pdf_bytes("Project Phoenix budget approved")
    doc = Document(owner_id=user.id, filename="memo.pdf", content_type="application/pdf",
                   size_bytes=len(raw), status=DocumentStatus.PENDING)
    session.add(doc)
    session.commit()
    ingest_document(session, doc.id, raw, embedder, store)
    session.refresh(doc)
    assert doc.status == DocumentStatus.READY
    assert doc.chunk_count >= 1
    chunks = store.search(session, user.id, embedder.embed_query("Phoenix budget"), top_k=3)
    assert any("Phoenix" in c.content for c in chunks)
    session.close()
